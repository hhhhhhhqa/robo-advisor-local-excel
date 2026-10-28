from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook

from robo_advisor.excel_frontend import TEMPLATE_PATH, run_excel_frontend
from robo_advisor.market_data import MarketData, load_fx_snapshot, load_price_snapshot


def _fixture(tmp_path: Path):
    input_path = tmp_path / "input.xlsx"
    output_path = input_path
    prices_path = tmp_path / "prices.parquet"
    fx_path = tmp_path / "fx.parquet"
    book = load_workbook(TEMPLATE_PATH)
    book["Dashboard"]["B4"] = 1000
    book["Dashboard"]["E3"] = datetime(2026, 9, 4)
    trades = book["Dashboard"]
    rows = [
        [datetime(2026, 9, 1), "Exchange", None, 500, "USD", "JPY"],
        [datetime(2026, 9, 1), "Buy", "SPY", 2, "JPY", None],
        [datetime(2026, 9, 3), "Sell", "SPY", 1, "JPY", None],
        [datetime(2026, 9, 4), "Exchange", None, 39000, "JPY", "USD"],
    ]
    for row_number, values in enumerate(rows, 27):
        for col, value in enumerate(values, 1):
            trades.cell(row_number, col, value)
    book.save(input_path)
    pd.DataFrame({
        "date": pd.to_datetime(["2026-09-01", "2026-09-02", "2026-09-03", "2026-09-04"]),
        "ticker": ["SPY"] * 4,
        "Close": [100, 110, 90, 120],
    }).to_parquet(prices_path)
    pd.DataFrame({
        "date": pd.to_datetime(["2026-09-01", "2026-09-02", "2026-09-03", "2026-09-04"]),
        "currency": ["JPY"] * 4,
        "usd_per_unit": [0.01, 0.009, 0.01, 0.01],
        "symbol": ["JPY=X"] * 4,
    }).to_parquet(fx_path)
    return input_path, output_path, prices_path, fx_path


def test_trades_fx_history_and_monitoring(tmp_path: Path):
    paths = _fixture(tmp_path)
    run_excel_frontend(*paths)
    result = load_workbook(paths[1], data_only=False)
    assert result.sheetnames == ["Dashboard"]
    sheet = result["Dashboard"]
    assert len(sheet._charts) == 1
    assert len(sheet._charts[0].series) == 2
    assert sheet["F8"].value == 4
    assert sheet["E8"].value == "=INDEX($A$440:$A$443,$F$8)"
    assert "INDEX" in sheet["E9"].value
    assert sheet["I443"].value == "=IF(ROW()-439=$F$8,$B443,NA())"
    assert sheet["J27"].value == 1
    assert [sheet.cell(row, 2).value for row in range(440, 444)] == pytest.approx(
        [1000, 990, 980, 1010]
    )
    assert sheet["E443"].value == pytest.approx(0.01)
    assert sheet["G443"].value == pytest.approx(-0.02)
    assert sheet["F443"].value is not None
    assert sheet["B231"].value == "Exchange"
    assert sheet["K232"].value == 20000
    assert sheet["M233"].value == 1
    with ZipFile(paths[1]) as archive:
        control = archive.read("xl/ctrlProps/ctrlProp1.xml").decode()
        assert 'fmlaLink="$F$8"' in control
        assert 'max="4"' in control

    run_excel_frontend(*paths)
    with ZipFile(paths[1]) as archive:
        assert archive.namelist().count("xl/ctrlProps/ctrlProp1.xml") == 1


def test_rejects_oversell_and_missing_fx(tmp_path: Path):
    paths = _fixture(tmp_path)
    book = load_workbook(paths[0])
    book["Dashboard"]["D29"] = 3
    book.save(paths[0])
    with pytest.raises(ValueError, match="insufficient SPY shares"):
        run_excel_frontend(*paths)

    book = load_workbook(paths[0])
    book["Dashboard"]["D29"] = 1
    book.save(paths[0])
    paths[3].unlink()
    with pytest.raises(ValueError, match="update FX rates"):
        run_excel_frontend(*paths)


def test_no_refresh_uses_cache_without_download(tmp_path: Path, monkeypatch):
    paths = _fixture(tmp_path)
    import robo_advisor.market_data as data
    monkeypatch.setattr(data.yf, "Ticker", lambda *args: pytest.fail("unexpected network request"))
    assert len(load_price_snapshot(paths[2], False, ["SPY"], "2020-01-01")) == 4
    assert len(load_fx_snapshot(paths[3], False, "2020-01-01")) == 4
    run_excel_frontend(*paths)


def test_rejects_trades_out_of_date_order(tmp_path: Path):
    paths = _fixture(tmp_path)
    book = load_workbook(paths[0])
    book["Dashboard"]["A27"] = datetime(2026, 9, 2)
    book.save(paths[0])
    with pytest.raises(ValueError, match="date order"):
        run_excel_frontend(*paths)


def test_trade_needs_its_own_etf_close_on_selected_day(tmp_path: Path):
    paths = _fixture(tmp_path)
    book = load_workbook(paths[0])
    book["Dashboard"]["A28"] = datetime(2026, 9, 2)
    book.save(paths[0])
    prices = pd.read_parquet(paths[2])
    prices = prices.loc[~((prices["ticker"] == "SPY") & (prices["date"] == pd.Timestamp("2026-09-02")))]
    prices = pd.concat([prices, pd.DataFrame({
        "date": [pd.Timestamp("2026-09-02")], "ticker": ["AGG"], "Close": [50]
    })], ignore_index=True)
    prices.to_parquet(paths[2])
    with pytest.raises(ValueError, match="No ETF close for SPY on 2026-09-02"):
        run_excel_frontend(*paths)


def test_selective_refresh_and_fx_quote_directions(tmp_path: Path, monkeypatch):
    import robo_advisor.market_data as data

    downloaded = []

    def price_download(ticker, start):
        downloaded.append((ticker, start))
        return pd.DataFrame({
            "date": pd.to_datetime(["2026-09-02"]),
            "ticker": [ticker],
            "Close": [120],
        }), "America/New_York"

    class FakeTicker:
        def __init__(self, symbol):
            downloaded.append(symbol)
            self.symbol = symbol

        def history(self, **kwargs):
            value = {"JPY=X": 100, "EURUSD=X": 1.2, "AUDUSD=X": 0.7, "CNY=X": 7}[self.symbol]
            return pd.DataFrame({"Close": [value]},
                                index=pd.DatetimeIndex(["2026-09-02"], tz="UTC"))

    monkeypatch.setattr(data, "_download_one", price_download)
    monkeypatch.setattr(data.yf, "Ticker", FakeTicker)
    prices_path, fx_path = tmp_path / "prices.parquet", tmp_path / "fx.parquet"
    pd.DataFrame({"date": pd.to_datetime(["2026-09-01"]), "ticker": ["SPY"],
                  "Close": [100]}).to_parquet(prices_path)
    prices = load_price_snapshot(prices_path, True, ["SPY"], "2020-01-01")
    fx = load_fx_snapshot(fx_path, True, "2020-01-01")
    market = MarketData(prices, fx)
    assert len(prices) == 2
    assert downloaded[0][0] == "SPY"
    assert market.usd_per_unit(date(2026, 9, 2), "JPY")[0] == pytest.approx(0.01)
    assert market.usd_per_unit(date(2026, 9, 2), "EUR")[0] == pytest.approx(1.2)
    assert market.usd_per_unit(date(2026, 9, 2), "AUD")[0] == pytest.approx(0.7)
    assert market.usd_per_unit(date(2026, 9, 2), "CNY")[0] == pytest.approx(1 / 7)
    assert len(load_price_snapshot(prices_path, False, ["SPY"], "2020-01-01")) == 2
    assert len(downloaded) == 5


def test_fx_must_be_observed_no_later_than_valuation_day():
    prices = pd.DataFrame({"date": pd.to_datetime(["2026-09-01"]), "ticker": ["SPY"], "Close": [100]})
    fx = pd.DataFrame({"date": pd.to_datetime(["2026-09-02"]), "currency": ["JPY"],
                       "usd_per_unit": [0.01], "symbol": ["JPY=X"]})
    market = MarketData(prices, fx)
    with pytest.raises(ValueError, match="on or before"):
        market.usd_per_unit(date(2026, 9, 1), "JPY")
