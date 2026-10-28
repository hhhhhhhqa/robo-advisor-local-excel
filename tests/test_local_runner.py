"""Exercise distribution isolation and no-overwrite guarantees without network."""
from datetime import datetime
import importlib.util
from pathlib import Path
import shutil
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook

from robo_advisor.excel_frontend import TEMPLATE_PATH

spec = importlib.util.spec_from_file_location("local_runner", Path(__file__).resolve().parents[1] / "local_runner.py")
runner = importlib.util.module_from_spec(spec)
spec.loader.exec_module(runner)


def local_folder(root: Path, starting_cash=1000):
    (root / "examples" / "excel").mkdir(parents=True)
    shutil.copy2(TEMPLATE_PATH, root / "examples" / "excel" / TEMPLATE_PATH.name)
    (root / "config").mkdir()
    shutil.copy2(TEMPLATE_PATH.parents[2] / "config" / "assets.json", root / "config" / "assets.json")
    workbook = runner.prepare(root)
    book = load_workbook(workbook)
    sheet = book["Dashboard"]
    sheet["B4"] = starting_cash
    sheet["E3"] = datetime(2026, 9, 4)
    for row, values in enumerate([
        [datetime(2026, 9, 1), "Exchange", None, 500, "USD", "JPY"],
        [datetime(2026, 9, 1), "Buy", "SPY", 2, "JPY", None],
        [datetime(2026, 9, 3), "Sell", "SPY", 1, "JPY", None],
        [datetime(2026, 9, 4), "Exchange", None, 39000, "JPY", "USD"],
    ], 27):
        for column, value in enumerate(values, 1):
            sheet.cell(row, column, value)
    book.save(workbook)
    market = root / "local_data" / "market"
    market.mkdir(parents=True)
    days = pd.to_datetime(["2026-09-01", "2026-09-02", "2026-09-03", "2026-09-04"])
    pd.DataFrame({"date": days, "ticker": ["SPY"] * 4,
                  "Close": [100, 110, 90, 120]}).to_parquet(market / "etf_prices.parquet")
    pd.DataFrame({"date": days, "currency": ["JPY"] * 4,
                  "usd_per_unit": [0.01, 0.009, 0.01, 0.01],
                  "symbol": ["JPY=X"] * 4}).to_parquet(market / "fx_rates.parquet")
    return workbook


@pytest.fixture(autouse=True)
def no_unexpected_network(monkeypatch):
    import robo_advisor.market_data as market
    monkeypatch.setattr(market.yf, "Ticker", lambda *a, **kw: pytest.fail("unexpected network request"))


def test_two_downloads_are_independent_and_retain_original_workflow(tmp_path):
    left, right = tmp_path / "Alice's folder 空格", tmp_path / "Bob"
    left_book = local_folder(left)
    right_book = local_folder(right, starting_cash=2000)
    right_before = right_book.read_bytes()
    left_before = left_book.read_bytes()
    with runner.folder_lock(left):
        runner.update(left, refresh=False)
    assert right_book.read_bytes() == right_before
    sheet = load_workbook(left_book)["Dashboard"]
    assert [sheet.cell(row, 2).value for row in range(440, 444)] == [1000, 990, 980, 1010]
    assert sheet["A27"].value == datetime(2026, 9, 1)
    assert sheet["E443"].value == pytest.approx(.01)
    assert len(sheet._charts) == 1
    with ZipFile(left_book) as archive:
        assert 'fmlaLink="$F$8"' in archive.read("xl/ctrlProps/ctrlProp1.xml").decode()
    backups = list((left / "local_data" / "backups").glob("*.xlsx"))
    assert len(backups) == 1
    assert backups[0].read_bytes() == left_before
    runner.update(right, refresh=False)
    assert load_workbook(right_book)["Dashboard"]["B443"].value == 2010
    assert load_workbook(left_book)["Dashboard"]["B443"].value == 1010


def test_reinitialization_never_resets_personal_trades(tmp_path):
    workbook = local_folder(tmp_path / "person")
    before = workbook.read_bytes()
    assert runner.prepare(workbook.parent) == workbook
    assert workbook.read_bytes() == before


def test_invalid_trade_leaves_workbook_unchanged(tmp_path):
    workbook = local_folder(tmp_path / "person")
    book = load_workbook(workbook)
    book["Dashboard"]["D29"] = 3
    book.save(workbook)
    before = workbook.read_bytes()
    with pytest.raises(ValueError, match="insufficient SPY shares"):
        runner.update(workbook.parent, refresh=False)
    assert workbook.read_bytes() == before
    assert not (workbook.parent / "local_data" / "backups").exists()


def test_incomplete_or_failed_download_never_overwrites_input(tmp_path, monkeypatch):
    import robo_advisor.market_data as market
    workbook = local_folder(tmp_path / "person")
    (workbook.parent / "local_data" / "market" / "etf_prices.parquet").unlink()
    before = workbook.read_bytes()
    def fail(*args, **kwargs):
        raise RuntimeError("network unavailable")
    monkeypatch.setattr(market, "load_price_snapshot", fail)
    with pytest.raises(RuntimeError, match="network unavailable"):
        runner.update(workbook.parent)
    assert workbook.read_bytes() == before


def test_open_excel_is_rejected_before_any_download(tmp_path):
    workbook = local_folder(tmp_path / "person")
    workbook.with_name("~$Portfolio.xlsx").touch()
    before = workbook.read_bytes()
    with pytest.raises(runner.LocalError, match="关闭"):
        runner.update(workbook.parent, refresh=False)
    assert workbook.read_bytes() == before


def test_concurrent_edit_is_not_overwritten(tmp_path, monkeypatch):
    import robo_advisor.excel_frontend as frontend
    workbook = local_folder(tmp_path / "person")
    write_results = frontend.write_results
    def edit_during_calculation(*args, **kwargs):
        write_results(*args, **kwargs)
        book = load_workbook(workbook)
        book["Dashboard"]["B4"] = 9000
        book.save(workbook)
    monkeypatch.setattr(frontend, "write_results", edit_during_calculation)
    with pytest.raises(runner.LocalError, match="发生变化"):
        runner.update(workbook.parent, refresh=False)
    assert load_workbook(workbook)["Dashboard"]["B4"].value == 9000


def test_lock_rejects_second_update_and_releases_after_error(tmp_path):
    with runner.folder_lock(tmp_path):
        with pytest.raises(runner.LocalError, match="正在更新"):
            with runner.folder_lock(tmp_path):
                pytest.fail("second update was allowed")
    with runner.folder_lock(tmp_path):
        pass


def test_manual_update_fetches_both_caches_despite_legacy_no_switches(tmp_path, monkeypatch):
    import robo_advisor.market_data as market
    workbook = local_folder(tmp_path / "person")
    market_dir = workbook.parent / "local_data" / "market"
    price_data = pd.read_parquet(market_dir / "etf_prices.parquet")
    fx_data = pd.read_parquet(market_dir / "fx_rates.parquet")
    shutil.rmtree(market_dir)
    requested = []
    def prices(path, refresh, tickers, start):
        requested.append((path, refresh))
        price_data.to_parquet(path)
        return price_data
    def fx(path, refresh, start):
        requested.append((path, refresh))
        fx_data.to_parquet(path)
        return fx_data
    monkeypatch.setattr(market, "load_price_snapshot", prices)
    monkeypatch.setattr(market, "load_fx_snapshot", fx)
    runner.update(workbook.parent)
    assert len(requested) == 2 and all(refresh for _, refresh in requested)
    assert all(path.exists() for path in market_dir.glob("*.parquet"))


def test_write_failure_retains_input_and_staging_is_cleaned(tmp_path, monkeypatch):
    import robo_advisor.excel_frontend as frontend
    workbook = local_folder(tmp_path / "person")
    before = workbook.read_bytes()
    def fail(*args, **kwargs):
        raise OSError("disk full")
    monkeypatch.setattr(frontend, "write_results", fail)
    with pytest.raises(OSError, match="disk full"):
        runner.update(workbook.parent, refresh=False)
    assert workbook.read_bytes() == before
    assert not list((workbook.parent / "local_data").glob("update-*"))


def test_automatic_calculation_never_downloads_even_with_legacy_yes_flags(tmp_path, monkeypatch):
    import socket
    import robo_advisor.market_data as market
    workbook = local_folder(tmp_path / "person")
    book = load_workbook(workbook)
    book["Dashboard"]["E5"] = "Yes"
    book["Dashboard"]["E6"] = "Yes"
    book["Dashboard"]["B4"] = 2000
    book.save(workbook)
    def forbidden(*args, **kwargs):
        pytest.fail("Automatic calculation attempted a network request")
    monkeypatch.setattr(socket.socket, "connect", forbidden)
    monkeypatch.setattr(market, "_download_one", forbidden)
    monkeypatch.setattr(market.yf, "Ticker", forbidden)
    result = workbook.parent / "local_data/result.xlsx"
    caches = {p: p.read_bytes() for p in (workbook.parent / "local_data/market").glob("*.parquet")}
    runner.calculate_for_excel(workbook.parent, workbook, result)
    assert load_workbook(result)["Dashboard"]["B443"].value == 2010
    assert all(p.read_bytes() == data for p, data in caches.items())
    # Editing a trade changes the historical path using the same cached prices.
    book["Dashboard"]["D29"] = 2
    book.save(workbook)
    runner.calculate_for_excel(workbook.parent, workbook, result)
    sheet = load_workbook(result)["Dashboard"]
    assert sheet["B443"].value == 1980
    assert sheet["J27"].value is None  # the position was fully sold
    assert [sheet.cell(r, 2).value for r in range(440, 444)] == [2000, 1990, 1980, 1980]


def test_missing_cache_requires_manual_update_without_any_download(tmp_path, monkeypatch):
    import robo_advisor.market_data as market
    workbook = local_folder(tmp_path / "person")
    shutil.rmtree(workbook.parent / "local_data/market")
    monkeypatch.setattr(market, "load_price_snapshot", lambda *a, **kw: pytest.fail("cache loader must not download"))
    before = workbook.read_bytes()
    with pytest.raises(runner.LocalError, match="Click Update"):
        runner.calculate_for_excel(workbook.parent, workbook, workbook.parent / "local_data/result.xlsx")
    assert workbook.read_bytes() == before


def test_partial_manual_download_keeps_both_previous_caches(tmp_path, monkeypatch):
    import robo_advisor.market_data as market
    workbook = local_folder(tmp_path / "person")
    directory = workbook.parent / "local_data/market"
    before = {p: p.read_bytes() for p in directory.glob("*.parquet")}
    def prices(path, *args):
        frame = pd.read_parquet(path)
        frame["Close"] = frame["Close"] * 2
        frame.to_parquet(path)
        return frame
    def fail(*args):
        raise RuntimeError("FX API unavailable")
    monkeypatch.setattr(market, "load_price_snapshot", prices)
    monkeypatch.setattr(market, "load_fx_snapshot", fail)
    original = workbook.read_bytes()
    with pytest.raises(RuntimeError, match="FX API unavailable"):
        runner.update(workbook.parent)
    assert all(p.read_bytes() == data for p, data in before.items())
    assert workbook.read_bytes() == original
