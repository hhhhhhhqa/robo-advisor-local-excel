from pathlib import Path
import shutil
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from test_local_runner import local_folder, runner
from workbook_button import attach_project, prepare_button_workbook, template_path

PROJECT = Path(__file__).resolve().parents[1]


def add_template(root):
    target = template_path(root)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path(PROJECT), target)


def test_native_template_has_button_and_no_local_folder_paths():
    with ZipFile(template_path(PROJECT)) as archive:
        assert archive.read("xl/vbaProject.bin").startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
        assert b"Update" in archive.read("xl/drawings/drawing1.xml")
        for name in archive.namelist():
            if name.endswith(".xml"):
                assert b"/Users/" not in archive.read(name)


def test_embedded_macros_match_reviewable_source_and_use_workbook_events():
    from oletools.olevba import VBA_Parser
    parser = VBA_Parser(str(template_path(PROJECT)))
    embedded = {name: code for _, _, name, code in parser.extract_macros()}
    parser.close()
    for filename in ("Module1.bas", "ThisWorkbook.cls"):
        expected = (PROJECT / "excel_button" / filename).read_text().strip()
        assert embedded[filename].replace("\r\n", "\n").strip() == expected
    assert "Workbook_SheetChange" in embedded["ThisWorkbook.cls"]
    assert "RunPortfolio False" in embedded["Module1.bas"]
    assert "RunPortfolio True" in embedded["Module1.bas"]


def test_upgrade_preserves_existing_inputs_and_the_legacy_file(tmp_path):
    root = tmp_path / "person"
    legacy = local_folder(root)
    add_template(root)
    before = legacy.read_bytes()
    workbook = prepare_button_workbook(root)
    assert workbook.name == "Portfolio.xlsm"
    assert legacy.read_bytes() == before
    book = load_workbook(workbook)
    assert book["Dashboard"]["B4"].value == 1000
    assert book["Dashboard"]["C28"].value == "SPY"
    with ZipFile(workbook) as archive, ZipFile(template_path(PROJECT)) as template:
        assert archive.read("xl/vbaProject.bin") == template.read("xl/vbaProject.bin")
        assert b"macroEnabled.main+xml" in archive.read("[Content_Types].xml")
        for name in ("xl/workbook.xml", "xl/worksheets/sheet1.xml"):
            ET.fromstring(archive.read(name))
    saved = workbook.read_bytes()
    prepare_button_workbook(root)
    assert workbook.read_bytes() == saved


def test_button_calculates_snapshot_without_replacing_open_workbook(tmp_path, monkeypatch):
    import robo_advisor.market_data as market
    monkeypatch.setattr(market.yf, "Ticker", lambda *a, **kw: pytest.fail("unexpected network request"))
    root = tmp_path / "person"
    legacy = local_folder(root)
    add_template(root)
    workbook = prepare_button_workbook(root)
    workbook.with_name("~$Portfolio.xlsm").touch()
    before = workbook.read_bytes()
    snapshot = root / "local_data" / "button_input.xlsm"
    shutil.copy2(workbook, snapshot)
    result = runner.calculate_for_excel(root)
    assert workbook.read_bytes() == before
    sheet = load_workbook(result)["Dashboard"]
    assert sheet["B443"].value == 1010
    assert sheet["G443"].value == pytest.approx(-.02)
    assert len(sheet._charts) == 1
    with ZipFile(result) as archive:
        assert "xl/vbaProject.bin" not in archive.namelist()
        assert "xl/ctrlProps/ctrlProp1.xml" in archive.namelist()
    backup = list((root / "local_data" / "backups").glob("*.xlsm"))[0]
    assert backup.read_bytes() == snapshot.read_bytes()


def test_failed_button_run_removes_stale_result_and_preserves_input(tmp_path):
    root = tmp_path / "person"
    workbook = local_folder(root)
    snapshot = root / "local_data" / "button_input.xlsm"
    shutil.copy2(workbook, snapshot)
    book = load_workbook(snapshot)
    book["Dashboard"]["D29"] = 3
    book.save(snapshot)
    output = root / "local_data" / "button_result.xlsx"
    shutil.copy2(workbook, output)
    before = workbook.read_bytes()
    with pytest.raises(ValueError, match="insufficient SPY shares"):
        runner.calculate_for_excel(root)
    assert not output.exists()
    assert workbook.read_bytes() == before


def test_mac_style_temporary_paths_and_external_update_keep_macros(tmp_path, monkeypatch):
    import robo_advisor.market_data as market
    monkeypatch.setattr(market.yf, "Ticker", lambda *a, **kw: pytest.fail("unexpected network request"))
    root = tmp_path / "person"
    local_folder(root)
    add_template(root)
    workbook = prepare_button_workbook(root)
    sandbox = tmp_path / "Excel temporary 空格's"
    sandbox.mkdir()
    snapshot = sandbox / "input.xlsm"
    output = sandbox / "result.xlsx"
    shutil.copy2(workbook, snapshot)
    runner.calculate_for_excel(root, snapshot, output)
    assert load_workbook(output)["Dashboard"]["B443"].value == 1010
    runner.update(root, refresh=False)
    with ZipFile(workbook) as archive:
        assert "xl/vbaProject.bin" in archive.namelist()
        assert "xl/ctrlProps/ctrlProp1.xml" in archive.namelist()
    assert load_workbook(workbook)["Dashboard"]["B443"].value == 1010


def test_button_output_cannot_replace_its_input(tmp_path):
    root = tmp_path / "person"
    workbook = local_folder(root)
    before = workbook.read_bytes()
    with pytest.raises(runner.LocalError, match="分开保存"):
        runner.calculate_for_excel(root, workbook, workbook)
    assert workbook.read_bytes() == before
